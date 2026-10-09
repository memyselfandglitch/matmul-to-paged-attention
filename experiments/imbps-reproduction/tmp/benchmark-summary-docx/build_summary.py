from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/deveshi/study/experiments/imbps-reproduction")
OUT = ROOT / "output/CPU_Inference_Benchmarking_Results_Summary.docx"
TMP = ROOT / "tmp/benchmark-summary-docx"
CHART = TMP / "decode_crossover.png"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "5B6573"
LIGHT_GRAY = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
GREEN = "2E7D32"
RED = "A63D40"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=105, bottom=90, end=105) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold=False, color=BLACK, size=8.7,
                  align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Liberation Sans"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def add_table(doc, headers, rows, widths=None, alignments=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        set_cell_text(hdr.cells[i], header, bold=True, color=WHITE, size=font_size,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i], NAVY)
        if widths:
            hdr.cells[i].width = widths[i]
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            align = alignments[ci] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[ci], value, size=font_size, align=align)
            set_cell_shading(cells[ci], PALE_BLUE if ri % 2 else WHITE)
            if widths:
                cells[ci].width = widths[ci]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_bullet(doc, text: str, *, color=BLACK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def make_chart():
    path = ROOT / "results/decode-l2-pilot-9970/timing/summary.csv"
    with path.open() as f:
        rows = list(csv.DictReader(f))
    k4 = sorted(
        [r for r in rows if r["backend"] == "imbps" and r["split"] == "4"],
        key=lambda r: int(r["batch"]),
    )
    x = [int(r["batch"]) for r in k4]
    y = [float(r["paired_speedup_median"]) for r in k4]
    lo = [float(r["paired_speedup_ci95_low"]) for r in k4]
    hi = [float(r["paired_speedup_ci95_high"]) for r in k4]
    yerr = [[v - l for v, l in zip(y, lo)], [h - v for v, h in zip(y, hi)]]

    width, height = 1400, 600
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
    bold_path = Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")
    if not font_path.exists():
        font_path = Path("/Library/Fonts/Arial.ttf")
    if not bold_path.exists():
        bold_path = font_path
    font = ImageFont.truetype(str(font_path), 25) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 22) if font_path.exists() else ImageFont.load_default()
    title_font = ImageFont.truetype(str(bold_path), 31) if bold_path.exists() else font

    left, right, top, bottom = 120, 1345, 88, 500
    ymin, ymax = 0.42, 1.45
    import math

    def xpix(value):
        lo_x, hi_x = math.log2(min(x)), math.log2(max(x))
        return left + (math.log2(value) - lo_x) / (hi_x - lo_x) * (right - left)

    def ypix(value):
        return bottom - (value - ymin) / (ymax - ymin) * (bottom - top)

    # Beneficial region and grid.
    draw.rectangle((xpix(420), top, right, bottom), fill="#EAF2F8")
    for tick in (0.5, 0.75, 1.0, 1.25):
        py = ypix(tick)
        draw.line((left, py, right, py), fill="#D9D9D9", width=2)
        draw.text((left - 18, py), f"{tick:.2f}", fill="#5B6573", font=small, anchor="rm")
    draw.line((left, top, left, bottom), fill="#5B6573", width=2)
    draw.line((left, bottom, right, bottom), fill="#5B6573", width=2)

    # TPP baseline as a dashed line.
    by = ypix(1.0)
    dash = 18
    cursor = left
    while cursor < right:
        draw.line((cursor, by, min(cursor + dash, right), by), fill="#5B6573", width=3)
        cursor += dash * 2

    points = [(xpix(xv), ypix(yv)) for xv, yv in zip(x, y)]
    draw.line(points, fill="#2F75B5", width=5)
    for xv, yv, low, high in zip(x, y, lo, hi):
        px, py = xpix(xv), ypix(yv)
        lo_py, hi_py = ypix(low), ypix(high)
        draw.line((px, hi_py, px, lo_py), fill="#2F75B5", width=3)
        draw.line((px - 8, hi_py, px + 8, hi_py), fill="#2F75B5", width=3)
        draw.line((px - 8, lo_py, px + 8, lo_py), fill="#2F75B5", width=3)
        draw.ellipse((px - 8, py - 8, px + 8, py + 8), fill="#2F75B5", outline="white", width=2)
        draw.text((px, bottom + 16), str(xv), fill="#303842", font=small, anchor="ma")

    draw.text((width // 2, 25), "Preliminary decode shaped MLP crossover", fill="#000000", font=title_font, anchor="ma")
    draw.text(((left + right) // 2, 565), "Active rows M", fill="#303842", font=font, anchor="mm")
    y_label = Image.new("RGBA", (340, 50), (255, 255, 255, 0))
    y_draw = ImageDraw.Draw(y_label)
    y_draw.text((170, 25), "Speedup over TPP", fill="#303842", font=font, anchor="mm")
    y_label = y_label.rotate(90, expand=True)
    img.paste(y_label, (18, 165), y_label)
    draw.text((xpix(650), top + 12), "IMBPS beneficial", fill="#2E7D32", font=small, anchor="ma")
    draw.line((left + 16, top + 28, left + 64, top + 28), fill="#2F75B5", width=5)
    draw.text((left + 76, top + 28), "IMBPS K=4", fill="#303842", font=small, anchor="lm")
    draw.text((right - 15, by - 10), "TPP baseline", fill="#5B6573", font=small, anchor="ra")
    img.save(CHART, dpi=(180, 180))


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Liberation Sans"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title.element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    for name, size, before, after in (
        ("Heading 1", 15, 13, 6),
        ("Heading 2", 11.5, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def main():
    TMP.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    make_chart()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    configure_styles(doc)

    # Footer
    footer = section.footer
    footer.distance = Cm(0.7)
    fp = footer.paragraphs[0]
    fp.text = "Page "
    fp.style = doc.styles["Normal"]
    fp.runs[0].font.size = Pt(7.5)
    fp.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_number(fp)

    # Page 1
    title = doc.add_paragraph("CPU Inference Benchmarking Results Summary", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Matrix multiplication  Paged KV cache  IMBPS")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    r.bold = True
    meta = doc.add_paragraph("Deveshi Singh   October 2026")
    meta.paragraph_format.space_after = Pt(16)
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(11)
    r = p.add_run("Main conclusion  ")
    r.bold = True
    p.add_run(
        "CPU inference performance depended on how work was grouped and how data was traversed. "
        "No single loop order or cache threshold predicted the fastest implementation. The strongest "
        "results came from reducing dispatch and allocation overhead, matching traversal to physical "
        "storage, and enabling IMBPS only when the active MLP workload was large enough."
    )

    add_heading(doc, "Results at a glance", 1)
    add_table(
        doc,
        ["Benchmark", "Observed result", "Practical interpretation"],
        [
            ["Batched matrix multiplication", "BMM remained faster after the nominal working set exceeded cache.", "Dispatches, temporary allocations and copies still mattered."],
            ["Optimized GEMM loop orders", "Six loop orders were usually within about 2 percent once they shared the same tiled SIMD microkernel.", "Microkernel quality mattered more than the outer loop order."],
            ["Paged KV cache", "Matched layout and traversal won; the preferred layout changed with physical block locality.", "Choose layout using the expected fragmentation pattern."],
            ["IMBPS", "Prefill gains were reproduced approximately. Decode shaped MLPs benefited only at high active row counts.", "Select K empirically and bypass IMBPS for small decode workloads."],
        ],
        widths=[Cm(4.0), Cm(6.5), Cm(6.7)],
        font_size=8.4,
    )

    add_heading(doc, "What changed across the experiments", 1)
    add_bullet(doc, "First, batched execution removed repeated framework dispatches and temporary result copies.")
    add_bullet(doc, "Next, KV experiments separated physical layout from traversal order and exposed a fragmentation dependent crossover.")
    add_bullet(doc, "Finally, IMBPS split the MLP expansion dimension to improve cache locality, but the best split still depended on kernel overhead and workload size.")

    page_break(doc)

    # Page 2
    add_heading(doc, "Matrix multiplication findings", 1)
    add_heading(doc, "Batched execution remained faster", 2)
    doc.add_paragraph(
        "A loop of individual matrix multiplications performed the same useful FLOPs as one batched "
        "matrix multiplication, but it issued one operation per batch item, created temporary outputs, "
        "allocated a final stacked tensor and copied each temporary into it. Those costs remained visible "
        "even after the matrices exceeded the nominal cache capacity."
    )
    add_table(
        doc,
        ["Implementation", "Execution and memory behavior"],
        [
            ["Looped MM", "Multiple aten::mm dispatches plus temporary C tensors and a final stacked C tensor"],
            ["Batched MM", "One aten::bmm dispatch and one final C allocation"],
        ],
        widths=[Cm(4.1), Cm(13.1)],
        font_size=8.7,
    )

    add_heading(doc, "The shared microkernel reduced loop order differences", 2)
    doc.add_paragraph(
        "Cache tiling plus SIMD register blocking was about 3.1 times faster than tiled-only multiplication "
        "in the earlier Apple ARM experiment. Once all six outer loop orders used that same microkernel, "
        "their geometric-mean performance was generally within about 2 percent of the best. The initial "
        "kij and ikj advantage was therefore workload-specific rather than universal."
    )

    add_heading(doc, "Paged KV cache findings", 1)
    doc.add_paragraph(
        "The KV experiments separated physical storage from traversal. Matching the traversal order to "
        "the physical layout consistently improved locality. Which matching layout was best depended on "
        "whether the physical block table contained long contiguous runs or highly fragmented blocks."
    )
    add_table(
        doc,
        ["Block table locality", "Preferred organization", "Observed result"],
        [
            ["Contiguous runs of at least 8 blocks", "HBND with head-first traversal", "At run length 512, HBND/BHND was 0.937 times"],
            ["Run length 4", "No clear winner", "Effective tie"],
            ["Runs of 1 or 2 blocks", "BHND with block-first traversal", "At run length 1, HBND/BHND was 1.059 times"],
        ],
        widths=[Cm(5.1), Cm(5.6), Cm(6.5)],
        font_size=8.5,
    )
    doc.add_paragraph(
        "A separate BF16 decode sweep found a repeatable region for block-major storage with block-first "
        "traversal. It won 39 of 48 MHA workloads with a median speedup of about 1.17 times, mainly at "
        "40 to 72 MiB of logical K plus V traffic per call. The baseline had no repeatable wins in that region."
    )

    page_break(doc)

    # Page 3
    add_heading(doc, "IMBPS prefill reproduction", 1)
    doc.add_paragraph(
        "IMBPS partitions the OPT-30B MLP expansion dimension and processes one block at a time. The "
        "reproduction used one EPYC 9654 socket with 96 physical cores and 384 MiB aggregate L3. The "
        "measured speedups were close to the paper, although K=8 slightly outperformed K=4 at the larger batches."
    )
    add_table(
        doc,
        ["Batch", "Paper K=4", "Measured K=4", "Fastest tested K"],
        [
            ["16", "1.21 times", "1.160 times", "K=4 at 1.160 times"],
            ["32", "1.25 times", "1.184 times", "K=8 at 1.185 times"],
            ["64", "1.23 times", "1.206 times", "K=8 at 1.213 times"],
        ],
        widths=[Cm(2.4), Cm(4.1), Cm(4.4), Cm(6.3)],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER] * 4,
        font_size=8.6,
    )

    add_heading(doc, "Cache fit did not guarantee speedup", 2)
    doc.add_paragraph(
        "For B=16 and sequence length 256, Equation 13 selected K=2 as the smallest power-of-two split "
        "whose modeled 364 MiB working set fit within 384 MiB L3. K=2 was statistically indistinguishable "
        "from TPP, and every larger split was slower. Cache residency alone was therefore not sufficient "
        "to predict runtime."
    )
    add_table(
        doc,
        ["Backend", "K", "Modeled working set", "Median time", "Speedup over TPP"],
        [
            ["TPP", "1", "672 MiB", "292.02 ms", "1.000 times"],
            ["IMBPS", "2", "364 MiB", "292.40 ms", "0.998 times  CI 0.978 to 1.046"],
            ["IMBPS", "4", "210 MiB", "454.19 ms", "0.645 times"],
            ["IMBPS", "8", "133 MiB", "330.55 ms", "0.883 times"],
            ["IMBPS", "16", "94.5 MiB", "358.31 ms", "0.818 times"],
            ["IMBPS", "32", "75.25 MiB", "613.19 ms", "0.478 times"],
        ],
        widths=[Cm(2.5), Cm(1.4), Cm(4.0), Cm(3.5), Cm(5.8)],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=8.1,
    )

    add_heading(doc, "The equation predicted locality better than runtime", 2)
    doc.add_paragraph(
        "At sequence lengths 1024 and 1408, the Equation 13 candidate uniquely minimized L3 misses and "
        "DRAM traffic and maximized hit rate and measured arithmetic intensity. The next larger split had "
        "the lowest profiled runtime. This supports using the equation as a starting point followed by a "
        "small empirical search."
    )
    add_table(
        doc,
        ["Sequence", "Equation K", "Best locality", "Lowest profiled time"],
        [
            ["1024", "16", "K=16", "K=32"],
            ["1408", "32", "K=32", "K=64"],
        ],
        widths=[Cm(3.0), Cm(3.2), Cm(5.2), Cm(5.8)],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER] * 4,
        font_size=8.7,
    )

    page_break(doc)

    # Page 4
    add_heading(doc, "Decode shaped MLP pilot", 1)
    doc.add_paragraph(
        "The current extension evaluates the standalone OPT-30B MLP with one token per sequence. Active "
        "rows M represent the live decode batch; for speculative verification, the same M can represent "
        "batch multiplied by verified draft tokens. The pilot does not include attention, KV-cache access, "
        "proposal generation or token acceptance."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(CHART), width=Cm(16.8))
    cap = doc.add_paragraph("Figure 1  K=4 became beneficial between 128 and 512 active rows")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)

    add_table(
        doc,
        ["Active rows", "TPP", "Best IMBPS K", "IMBPS time", "Paired speedup"],
        [
            ["1", "1.711 ms", "4", "3.419 ms", "0.501 times"],
            ["32", "3.591 ms", "4", "5.936 ms", "0.609 times"],
            ["128", "8.236 ms", "4", "14.201 ms", "0.581 times"],
            ["512", "39.001 ms", "4", "32.727 ms", "1.192 times  CI 1.182 to 1.198"],
            ["1024", "80.062 ms", "4", "58.420 ms", "1.370 times  CI 1.353 to 1.394"],
        ],
        widths=[Cm(2.7), Cm(3.1), Cm(3.1), Cm(3.2), Cm(5.1)],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=8.3,
    )
    doc.add_paragraph(
        "The pilot supports a conditional policy: use TPP for small active-row counts and consider IMBPS "
        "for highly batched decode or speculative verification. The full timing matrix is needed to locate "
        "the crossover more precisely. Counter evidence and end-to-end accepted-token throughput remain open."
    )

    page_break(doc)

    # Page 5
    add_heading(doc, "Experimental conditions", 1)
    add_table(
        doc,
        ["Item", "Earlier prefill experiments", "Decode pilot"],
        [
            ["Processor placement", "EPYC 9654 socket 0  96 cores  12 CCDs", "Same"],
            ["Cache model", "384 MiB aggregate L3", "96 MiB aggregate private L2"],
            ["Model operator", "OPT-30B MLP  BF16  ReLU  bias", "Same"],
            ["Software", "PACE 1.0.0  PyTorch 2.7.0 CPU", "Same"],
            ["Memory settings", "THP madvise  no tcmalloc preload", "THP always  tcmalloc 2.18.1"],
            ["Scheduling", "Exclusive Slurm allocation  one NUMA node", "Same"],
            ["Status", "Completed timing and historical uProf mechanism runs", "Preliminary three-round timing pilot"],
        ],
        widths=[Cm(3.7), Cm(7.0), Cm(6.5)],
        font_size=8.3,
    )

    add_heading(doc, "What is established and what remains open", 1)
    add_bullet(doc, "Established  Batched dispatch reduces framework and allocation overhead beyond what a cache-size-only hypothesis predicts.", color=GREEN)
    add_bullet(doc, "Established  KV layout should match traversal, and fragmentation changes the preferred layout.", color=GREEN)
    add_bullet(doc, "Established  IMBPS reproduced substantial OPT-30B prefill gains on mn01, although the optimal K differed from the paper at larger batches.", color=GREEN)
    add_bullet(doc, "Preliminary  Standalone decode-shaped MLPs benefited at 512 and 1024 active rows but regressed at 128 rows and below.", color=BLUE)
    add_bullet(doc, "Not yet established  End-to-end decode or speculative-decode throughput, because attention, KV-cache work and acceptance behavior were excluded.", color=RED)
    add_bullet(doc, "Not yet established  Parallel per-CCD IMBPS. Job 9300 only varied complete-CCD thread groups; it did not execute independent per-CCD splits.", color=RED)

    add_heading(doc, "Next experiments", 1)
    add_table(
        doc,
        ["Priority", "Experiment", "Decision it will support"],
        [
            ["1", "Complete the full decode timing matrix over M=1 to 1024 and K=2 to 32", "Choose a reliable active-row crossover and split policy"],
            ["2", "Add fixed-trace end-to-end normal and speculative decode", "Test accepted-token throughput with attention and KV cache included"],
            ["3", "Implement parallel IMBPS with one worker group per CCD", "Compare global splitting with CCD-local cache ownership and scaling"],
            ["4", "Repeat cache and traffic counters when amd_uncore access returns", "Confirm whether L2 and DRAM traffic explain the decode crossover"],
        ],
        widths=[Cm(1.6), Cm(8.2), Cm(7.4)],
        font_size=8.3,
    )

    add_heading(doc, "Evidence", 1)
    links = [
        ("Table II timing summary", "https://github.com/memyselfandglitch/matmul-to-paged-attention/blob/main/experiments/imbps-reproduction/results/table_ii-20261001T131423Z/summary.csv"),
        ("Cache fit control", "https://github.com/memyselfandglitch/matmul-to-paged-attention/blob/main/experiments/imbps-reproduction/results/cache_fit_opt30b-20261002T121900Z/summary.csv"),
        ("Equation 13 uProf summary", "https://github.com/memyselfandglitch/matmul-to-paged-attention/blob/main/experiments/imbps-reproduction/results/uprof-server-equation-opt30b-9392-analysis-9393/summary.csv"),
        ("Decode pilot summary", "https://github.com/memyselfandglitch/matmul-to-paged-attention/blob/main/experiments/imbps-reproduction/results/decode-l2-pilot-9970/timing/summary.csv"),
        ("Decode experimental plan", "https://github.com/memyselfandglitch/matmul-to-paged-attention/blob/main/experiments/imbps-reproduction/DECODE_L2_PLAN.md"),
    ]
    for label, url in links:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(2)
        add_hyperlink(p, label, url)

    # Keep document properties minimal and explicit.
    doc.core_properties.title = "CPU Inference Benchmarking Results Summary"
    doc.core_properties.subject = "Matrix multiplication paged KV cache and IMBPS benchmark findings"
    doc.core_properties.author = "Deveshi Singh"
    doc.core_properties.keywords = "CPU inference, KV cache, IMBPS, AMD EPYC, benchmarking"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
